"""Create the editable Word sample, then insert the supplied SVGs as OOXML images."""
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

TITLE = "000099"
FOCUS = "28317B"
COMPARE = "7F7F7F"
BODY = "111111"
SECONDARY = "4A4A4A"
NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def run(text, size=11, color=BODY, bold=False):
    text.font.name = "Microsoft JhengHei"
    text._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    text.font.size = Pt(size)
    text.font.bold = bold
    text.font.color.rgb = RGBColor.from_string(color)


def run_heading(text, size=20):
    text.font.name = "Cambria"
    text._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
    text._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
    text._element.rPr.rFonts.set(qn("w:eastAsia"), "PMingLiU")
    text.font.size = Pt(size)
    text.font.bold = True
    text.font.color.rgb = RGBColor.from_string(TITLE)


def drawing(rid, name, docpr_id):
    return f'''<w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="5334000" cy="2333625"/><wp:docPr id="{docpr_id}" name="{name}" descr="{name}"/><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="{name}"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="5334000" cy="2333625"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'''


def inject_svgs(docx_path, before_svg, after_svg):
    replacements = {
        "word/media/chart-before.svg": before_svg.read_bytes(),
        "word/media/chart-after.svg": after_svg.read_bytes(),
    }
    with zipfile.ZipFile(docx_path, "r") as source:
        files = {name: source.read(name) for name in source.namelist()}
    document = files["word/document.xml"].decode("utf-8")
    document = document.replace('<w:document ', '<w:document xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" ', 1)
    document = re.sub(r"<w:r>(?:(?!</w:r>).)*?<w:t>\[SVG_BEFORE\]</w:t></w:r>", drawing("rIdSvgBefore", "chart-before.svg", 1001), document, count=1)
    document = re.sub(r"<w:r>(?:(?!</w:r>).)*?<w:t>\[SVG_AFTER\]</w:t></w:r>", drawing("rIdSvgAfter", "chart-after.svg", 1002), document, count=1)
    files["word/document.xml"] = document.encode("utf-8")
    rels = files["word/_rels/document.xml.rels"].decode("utf-8")
    rels = rels.replace("</Relationships>", '<Relationship Id="rIdSvgBefore" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/chart-before.svg"/><Relationship Id="rIdSvgAfter" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/chart-after.svg"/></Relationships>')
    files["word/_rels/document.xml.rels"] = rels.encode("utf-8")
    types = files["[Content_Types].xml"].decode("utf-8")
    if 'Extension="svg"' not in types:
        types = types.replace("</Types>", '<Default Extension="svg" ContentType="image/svg+xml"/></Types>')
    files["[Content_Types].xml"] = types.encode("utf-8")
    files.update(replacements)
    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in files.items():
            target.writestr(name, content)


def main(data_path, before_path, after_path, output_path):
    data = json.loads(Path(data_path).read_text())
    channels = data if isinstance(data, list) else data.get("channels", data.get("rows"))
    rows = [(item.get("name", item.get("channel")), item.get("value", item.get("amount"))) for item in channels]
    if sum(value for _, value in rows) != 300:
        raise ValueError("Illustrative data must total 300")
    doc = Document()
    for index, (label, accent, title) in enumerate((("Before", COMPARE, data.get("title_before", "通路保費組成")), ("After", FOCUS, data.get("title_after", "通路保費組成")))):
        heading = doc.add_heading(title, level=1)
        run_heading(heading.runs[0])
        paragraph = doc.add_paragraph()
        run(paragraph.add_run(label), 14, accent, True)
        period_p = doc.add_paragraph()
        run(period_p.add_run(data.get("period", "同一期間")), 11, SECONDARY)
        chart = doc.add_paragraph()
        run(chart.add_run("[SVG_BEFORE]" if index == 0 else "[SVG_AFTER]"))
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ("通路", "保費（百萬元）")):
            run(cell.paragraphs[0].add_run(value), 10, BODY, True)
        for name, value in rows + [("合計", 300)]:
            cells = table.add_row().cells
            run(cells[0].paragraphs[0].add_run(name), 10)
            run(cells[1].paragraphs[0].add_run(str(value)), 10)
        for body in data.get("paragraphs", []):
            summary = doc.add_paragraph()
            run(summary.add_run(body), 10, BODY)
        footnote = doc.add_paragraph()
        run(footnote.add_run(f"資料來源：{data.get('source', '校準用示意資料')}。"), 9, SECONDARY)
        if index == 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(output_path)
    inject_svgs(Path(output_path), Path(before_path), Path(after_path))


if __name__ == "__main__":
    main(*sys.argv[1:])
